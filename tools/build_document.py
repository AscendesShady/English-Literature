import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
stories=sorted([s for s in json.loads((ROOT/'content/catalog.json').read_text(encoding='utf-8')) if isinstance(s['source'],str)],key=lambda s:s['year'])
doc=Document()
section=doc.sections[0]
section.page_width=Inches(8.27);section.page_height=Inches(11.69)
section.top_margin=section.bottom_margin=Inches(.6)
section.left_margin=section.right_margin=Inches(.65)
for name in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
    style=doc.styles[name];style.font.name='Calibri'
    style.font.color.rgb=RGBColor.from_string('283C32' if name!='Title' else '000000')
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.space_after=Pt(3)
doc.styles['Normal'].paragraph_format.line_spacing=1
doc.styles['Title'].font.size=Pt(28)
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.styles['Subtitle'].font.color.rgb=RGBColor.from_string('6A7167')
doc.styles['Heading 1'].font.size=Pt(19)
doc.styles['Heading 2'].font.size=Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(6)
doc.styles['Heading 2'].paragraph_format.space_after=Pt(3)
doc.core_properties.title='Literature by Codex'
doc.core_properties.author='Codex'
doc.add_paragraph('Literature by Codex','Title')
doc.add_paragraph('73 additions with twelve narrated story cards','Subtitle')
doc.add_paragraph('This companion preserves all 73 works added by Codex: the original Odyssey, Antigone and Frankenstein additions, plus 70 further works from foundational novelists, dramatists and poets. Every entry includes the same twelve rewritten story cards, characters, memory cues, recall question and memorable quotation as the website. The narration begins by setting the scene, introduces people and motives as they matter, and follows choices into consequences. Read aloud or turn through the cards online to imagine each moment.')
doc.add_paragraph('Summaries and memory connections are original learning adaptations with plot spoilers. Poetry uses thematic and argumentative landmarks rather than invented plots. Short quotations follow the linked edition; translations and punctuation can vary. Publication, composition and performance dates are distinguished where necessary. Era groups overlap and serve as navigation aids.')
doc.add_paragraph('Your original three literature documents remain the source of the other 25 works and have not been edited. The website contains 98 works and 1176 scenes in total. Its installed app saves the whole library for offline reading after the first successful online visit.')
doc.add_paragraph('A reading path','Heading 2')
for author in ['William Shakespeare','William Wordsworth','Percy Bysshe Shelley','Jane Austen','Charles Dickens','George Eliot','T. S. Eliot','Virginia Woolf']:
    entries=[s['title'] for s in stories if author in s['author']]
    doc.add_paragraph(author+' — '+', '.join(entries))

def source_label(url):
    from urllib.parse import urlparse
    host=urlparse(url).hostname
    return {'www.gutenberg.org':'Project Gutenberg','www.gutenberg.net.au':'Project Gutenberg Australia','www.poetryfoundation.org':'Poetry Foundation','tseliot.com':'T. S. Eliot official website','viscomi.sites.oasis.unc.edu':'University of North Carolina'}.get(host,'Quotation reference')

def link(paragraph,label,url):
    rel=paragraph.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    el=OxmlElement('w:hyperlink');el.set(qn('r:id'),rel)
    run=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'355660');pr.append(color)
    underline=OxmlElement('w:u');underline.set(qn('w:val'),'single');pr.append(underline);run.append(pr)
    text=OxmlElement('w:t');text.text=label;run.append(text);el.append(run);paragraph._p.append(el)

for i,s in enumerate(stories):
    heading=doc.add_paragraph(s['title'],'Heading 1')
    heading.paragraph_format.page_break_before=True
    p=doc.add_paragraph();p.add_run(s['author']).bold=True;p.add_run('  |  '+s['date']+'  |  '+s['genre'])
    doc.add_paragraph('Setting: '+s['setting'])
    p=doc.add_paragraph(s['hook']);p.runs[0].italic=True
    doc.add_paragraph('Characters to remember','Heading 2')
    p=doc.add_paragraph()
    for j,c in enumerate(s['characters']):
        if j:p.add_run(' · ')
        p.add_run(c['name']+': ').bold=True;p.add_run(c['role'])
    doc.add_paragraph('The story in '+str(len(s['scenes']))+' scenes','Heading 2')
    for j,f in enumerate(s['scenes'],1):
        if j==7:
            continued=doc.add_paragraph(s['title']+' continued','Heading 2')
            continued.paragraph_format.page_break_before=True
            doc.add_paragraph('Scenes seven to twelve','Heading 3')
        p=doc.add_paragraph();p.paragraph_format.keep_with_next=True;p.add_run(str(j)+'  '+f['title']+' — ').bold=True
        p.add_run(f['text'])
        p=doc.add_paragraph('Memory cue: '+f['cue']);p.paragraph_format.space_after=Pt(4);p.runs[0].italic=True;p.runs[0].font.size=Pt(9)
    doc.add_paragraph('Recall and context','Heading 2')
    p=doc.add_paragraph();p.add_run(s['question']).bold=True
    doc.add_paragraph(s['answer'])
    doc.add_paragraph('Memorable quote','Heading 2')
    for q in s['quotes']:
        p=doc.add_paragraph('“'+q['text']+'”');p.runs[0].italic=True
        doc.add_paragraph(q['context'])
        doc.add_paragraph('Memory connection: '+q['meaning'])
        p=doc.add_paragraph('Quotation reference: ');link(p,source_label(q['source']),q['source'])
    if s['note']:doc.add_paragraph('Context: '+s['note'])
    p=doc.add_paragraph('Read the text: ');link(p,source_label(s['source']),s['source'])

footer=section.footer.paragraphs[0]
footer.alignment=2
footer.add_run('Literature by Codex  ·  ').font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

doc.save(ROOT/'literature by codex.docx')
print('Created literature by codex.docx with '+str(len(stories))+' works')
